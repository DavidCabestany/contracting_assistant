import time
import uuid
import httpx
import openpyxl
import os
from dotenv import load_dotenv
from tqdm import tqdm


load_dotenv(override=True)

API_URL = os.getenv("API_URL", "")
AUTH_TOKEN = os.getenv("AUTH_TOKEN", "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJzdWIiOiJ1c2VyLVVJIiwiZXhwIjoxNzcxMDYxNzQ2fQ.QU4NcRu4YGuJC0VA8mMH0qO92K8KeSyRjjAZ-0B-Jyo")
SLEEP_SECONDS = 5
HAIKU = os.getenv("HAIKU")
SONNET_V1 = os.getenv("SONNET_V1")
SONNET_V2 = os.getenv("SONNET_V2")
SONNET_37 = os.getenv("SONNET_37")
SONNET_45 = os.getenv("SONNET_45")
OPUS_3_V1 = os.getenv("OPUS_3_V1")

MODEL_NAME = "SONNET_45"
MODEL_ID = SONNET_45
RUN_COUNT = 1

question_sets = {
    "General": [
        # "Who is the local lawyer for CAMCAR?",
        # "what are AZ's standard publication terms?",
        # "who is the AZ lawyer for Italy?",
        # "what template do I use for AZ to license software?",
        # "can I use the Terms and Conditions as a framework agreement. (Guidance says this contract cannot be used for multiple purchases).",
        # "when do I need to include a cyber security appendix with an MSA (Guidance explains the requirements).",
        # "are there any specific payment terms within the EU or do the same ones apply across all Member States?",
        # "what is CCT Team?",
        # "how is the length of the payment term determined?",
        # "how should I proceed if I need a signature on my agreement?",
        # "What can CCT do for me?",
        # "Where can I seek information or guidance about contract templates?",
        # 'what modifications I can make within the "Expectations of Third Parties" clause if the contract is concluded by a Supplier being one individual who has no employees?',
        # "what is the difference between a warranty and an undertaking?",
        # "can we agree to limit Supplier's liability for wilful misconduct?",
        # "what is a handbook?",
        # "what can I use CAN Handbook for",
        # "What is AZ's position in relation to the indemnification language that is acceptable from AZ's point of view in scenarios where AZ is the sponsor of a study?",
        # "What is AZ's position on choosing the State of California US as the seat of dispute resolution for a Master Services Agreement?",
        # "What is AZ's position when a supplier requests to be named on AZ's insurance policy?",
        # "At AZ, how long are CDAs expected to be in force for?",
        # "Is it okay to accept late payment terms proposed by a supplier?",
        # "Can we accept payment of interest as a penalty on overdue invoices?",
        # "Can CCT tell me who can sign my contract?",
        # "Can I use a Supplier's template for a non-disclosure agreement?",
        # "What is the meaning of Novation?",
        # "how do I pick the correct data protection terms?",
        # "can AZ agree to pay the costs of an audit?",
        "What should be considered when backdating an agreement?",
        "What should be considered when back dating an agreement?",
        "What if the Counterparty wants to reduce the term from 5 years to 2 years in a CDA?",
        "How do I extend the duration of an existing CDA?",
        "do supplier subcontractors have to be approved in advance by AZ in Master Service Agreement",
        "Who supports the US Data Security Program Appendix? How can I reach out to them?",
        "what should I do if a counterparty under a CDA does not want any audit rights included in the CDA",
        "How do I determine which CDA template should be used?",
        "Is the timeline for payment of invoices determined by the seat of the supplier or the AZ entity contracting with the supplier?",
        "When should I use a one-way CDA?",
        "I am drafting a CDA, where both AZ and the Supplier will be exchanging confidential information, which type of CDA should I use in this case?",
        "Which type of document is commonly used when AZ is contracting for the sharing of confidential information, and where can I find it?",
        "What should I do if a supplier insists on having the same termination rights as AZ in an MSA",
        "What can I do if a supplier does not want to agree to AZ's 'expectations of third parties' clause in an MSA",
        "Are non-solicitation terms legal in all jurisdictions",
        "Where can I locate the CDA guidance?",
        "where can i find the tempate for nda/cd",
        "What standard IPRs does AZ grant to Supplier?",
        "What is the wording listed in the CAN Handbook for the IPR paragraph of the MSA?",
        'What do I do if a supplier wants to add the phrase "unforeseeable at the moment of the Agreement execution” in the definition of the Force Majeure Event in an MSA',
        "What should I do if a recipient of AZ's confidential information under the terms of a CDA wants to keep a copy of AZ's confidential information after termination of the CDA AZ signed with the recipient",
        "What should I do if a supplier under an MSA argues that the monetary caps for Data Protection and Cyber Security breaches are too high and should be reduced to the standard cap in the MSA",
        "Which guidance document can help me draft and negotiate an MSA?",
        "need to ensure its not linking R&D Templates and only is working with the DP Scope templates",
        "What topics are listed in the table of contents in the Can Handbook?",
        "Explain direct vs. indirect damages in terms of limitation of liability",
        "What should I do if a supplier pushes back on the 'Compliance with AZ Directions (Services and Deliverables.3.5) in the MSA",
        "what are some of the best practices for working with universities on their templates",
 
    ],
    "Privacy": [
        # "Can I agree to changing Clause 1.4?",
        # "We need to confirm what kind of Exhibit we need to add to an Agreement. This is about Medical Communications, which involves publications. The vendor is an Institution, and they have a database of clinical trials and patient data.",
        # "Is anonymized data personal data?",
        # "what template should I use when both parties are controllers?",
        # "What about when I process data of Chinese individuals? Which template to use?",
        # "Can I agree to limiting the time for AZ to object to appointment of Subprocessor?",
        # "Can I agree to extend deadline for the Supplier to delete personal data after expiration of the agreement?",
        # "Can I agree to changing the definition to the following wording: “Anonymised Data” data which does not itself identify any individual and which will not allow any individual to be re-identified, whether through its combination with other data held by an authorised party or otherwise;",
        # "Which template clarifies what all the other templates govern and when they are used?",
        # "What kind of template should I use when clinical data is in scope? AZ sponsors the clinical trial, and we decide on how the sample will be used.",
        "Where AZ is the controller and the Supplier is the processor within what time frame must data be deleted?",
        "Who supports the US Data Security Program Appendix? How can I reach out to them?",
        'I got the following wording of Anonymised Data definition. Does it align with what AZ standard? “Anonymised Data” data which does not itself identify any individual and which will not allow any individual to be re-identified, whether through its combination with other data held by an authorised party or otherwise;',
        "In line with AZ standards please confirm when a personal data breach must be notified to AZ?",
        "Can you please confirm the definition of personal data breach",
        "What are the general supplier obligations in a controller processor relationship",
        "Can a 30-day notice period be added to the audit clause in the AZ controller to Supplier processor Data Protection Appendix",
        "Should the breach notification timeline be extended to 72 hours",
        "Does the controller to controller data protection appendix require technical and organizational measures?",
        "for priro wirrten auhtorizaiton how many days for te supply to notify az of an intended changes concerning addiiton or replace emt of subprocessors?",
 
    ],
}
 
 

EXCEL_FILE = "qna_model_comparison.xlsx"

def ask_question(question, knowledge_type, model_id):
    payload = {
        "apiKey": "",
        "user": {
            "id": "kvcn639",
            "sessionId": "",
            "language": "us-en",
            "platform": "web",
        },
        "query": {
            "text": question,
            "knowledgeType": knowledge_type,
            "transactionCount": 0,
            "files": [],
        }
    }
    headers = {
        "accept": "application/json",
        "Authorization": f"Bearer {AUTH_TOKEN}",
        "Content-Type": "application/json",
    }
    try:
        response = httpx.post(API_URL, json=payload, headers=headers, timeout=180)
        response.raise_for_status()
        return response.json()
    except Exception as e:
        print(f"Error for model {model_id}, question '{question}': {e}")
        return {}

def get_or_create_sheet():
    if os.path.exists(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "QnA Test Results"
        ws.append(["Set", "Question"])
    return wb, ws

def find_question_row(ws, set_name, question):
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[0].value == set_name and row[1].value == question:
            return row[0].row
    return None


def main():
    wb, ws = get_or_create_sheet()

    # Find where to add new columns
    existing_headers = [cell.value for cell in ws[1]]
    new_headers = []
    for i in range(1, RUN_COUNT + 1):
        col_name = f"{MODEL_NAME}" if RUN_COUNT == 1 else f"{MODEL_NAME}_{i}"
        new_headers.append(col_name)
        if col_name not in existing_headers:
            ws.cell(row=1, column=len(existing_headers) + 1, value=col_name)
            existing_headers.append(col_name)
        # add citation columns for each run
        cit_col_name = f"{col_name}_Citations"
        if cit_col_name not in existing_headers:
            ws.cell(row=1, column=len(existing_headers) + 1, value=cit_col_name)
            existing_headers.append(cit_col_name)

    # Calculate total iterations for tqdm
    total_questions = sum(len(qs) for qs in question_sets.values())
    total_iterations = total_questions * RUN_COUNT

    with tqdm(total=total_iterations, desc="Processing QnA", ncols=100) as pbar:
        for set_name, questions in question_sets.items():
            for q in questions:
                # Find or create the row for this question
                row_num = find_question_row(ws, set_name, q)
                if not row_num:
                    ws.append([set_name, q])
                    row_num = ws.max_row

                for i in range(1, RUN_COUNT + 1):
                    result = ask_question(q, set_name, MODEL_ID)

                    # extract answer
                    answer = result.get("result", {}).get("answer", {}).get("ans", "") or "[No answer returned or error]"

                    # extract citations
                    citations = result.get("result", {}).get("citations", []) or []
                    # stringify concisely: "fileName#pageNumber|fileName#pageNumber"
                    citations_str = "|".join(
                        f"{c.get('fileName','')}#{c.get('pageNumber','')}"
                        for c in citations
                    )

                    # write answer
                    col_name = f"{MODEL_NAME}" if RUN_COUNT == 1 else f"{MODEL_NAME}_{i}"
                    col_idx = existing_headers.index(col_name) + 1
                    ws.cell(row=row_num, column=col_idx, value=answer)

                    # write citations to paired column
                    cit_col_name = f"{col_name}_Citations"
                    cit_col_idx = existing_headers.index(cit_col_name) + 1
                    ws.cell(row=row_num, column=cit_col_idx, value=citations_str)

                    # console output
                    print(f"[{set_name}] Q: {q}\nA: {answer}\nModel: {col_name}")
                    if citations:
                        print("Citations:")
                        for c in citations:
                            print(f"  - {c.get('fileName','')} (page {c.get('pageNumber','')})")
                    else:
                        print("Citations: [none]")
                    print("-" * 60)

                    time.sleep(SLEEP_SECONDS)
                    pbar.update(1)

        wb.save(EXCEL_FILE)
        print(f"Results saved to {EXCEL_FILE}")

if __name__ == "__main__":
    main()