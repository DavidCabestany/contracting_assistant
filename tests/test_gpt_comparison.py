from __future__ import annotations

import os
import time
import statistics
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

from dotenv import load_dotenv
from openai import AzureOpenAI, APIStatusError
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def must_get(name: str) -> str:
    v = os.getenv(name)
    if not v:
        raise RuntimeError(f"Missing env var: {name}")
    return v


def parse_deployments() -> List[str]:
    multi = os.getenv("AZURE_OPENAI_DEPLOYMENTS", "").strip()
    if multi:
        return [d.strip() for d in multi.split(",") if d.strip()]
    return [must_get("AZURE_OPENAI_DEPLOYMENT").strip()]


def percentile(values: List[float], p: float) -> Optional[float]:
    if not values:
        return None
    xs = sorted(values)
    k = int(round((len(xs) - 1) * p))
    return xs[k]


def safe_attr(obj: Any, name: str) -> Any:
    return getattr(obj, name, None)


def extract_reasoning_tokens(resp: Any) -> Optional[int]:
    usage = safe_attr(resp, "usage")
    if not usage:
        return None
    ctd = safe_attr(usage, "completion_tokens_details")
    if not ctd:
        return None
    return safe_attr(ctd, "reasoning_tokens")


def extract_text(resp: Any) -> str:
    try:
        msg = resp.choices[0].message
        content = getattr(msg, "content", None)
        if isinstance(content, str):
            return content
    except Exception:
        pass
    return ""


def call_chat_with_temp_fallback(
    client: AzureOpenAI,
    deployment: str,
    messages: List[Dict[str, str]],
    temperature: float,
) -> Tuple[Any, str]:
    try:
        resp = client.chat.completions.create(
            model=deployment,
            messages=messages,
            temperature=temperature,
        )
        return resp, f"temp={temperature}"
    except APIStatusError as e:
        txt = getattr(e.response, "text", "") or ""
        # Deployments that only support default temperature reject non-default
        if e.status_code == 400 and '"param": "temperature"' in txt:
            resp = client.chat.completions.create(
                model=deployment,
                messages=messages,
            )
            return resp, "temp=default"
        raise


def write_excel(raw_rows: List[Dict[str, Any]], out_path: str) -> None:
    wb = Workbook()

    ws_raw = wb.active
    ws_raw.title = "RawRuns"

    raw_headers = [
        "timestamp_utc",
        "deployment",
        "trial",
        "ok",
        "mode",
        "latency_ms",
        "prompt_tokens",
        "completion_tokens",
        "reasoning_tokens",
        "total_tokens",
        "finish_reason",
        "visible_chars",
        "error_http_status",
        "error_message",
        "output_preview",
    ]
    ws_raw.append(raw_headers)

    for row in raw_rows:
        ws_raw.append([row.get(h) for h in raw_headers])

    for col_idx, header in enumerate(raw_headers, start=1):
        max_len = len(header)
        for cell in ws_raw[get_column_letter(col_idx)]:
            if cell.value is None:
                continue
            max_len = max(max_len, len(str(cell.value)))
        ws_raw.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 70)

    ws_sum = wb.create_sheet("Summary")
    sum_headers = [
        "deployment",
        "runs",
        "successes",
        "failures",
        "empty_outputs",
        "empty_output_rate",
        "latency_ms_mean",
        "latency_ms_median",
        "latency_ms_p90",
        "prompt_tokens_mean",
        "completion_tokens_mean",
        "reasoning_tokens_mean",
        "total_tokens_mean",
    ]
    ws_sum.append(sum_headers)

    deployments = sorted(set(r["deployment"] for r in raw_rows))
    for dep in deployments:
        dep_rows = [r for r in raw_rows if r["deployment"] == dep]
        ok_rows = [r for r in dep_rows if r["ok"]]
        lat = [r["latency_ms"] for r in ok_rows if r["latency_ms"] is not None]

        empty_outputs = sum(1 for r in ok_rows if (r.get("visible_chars") or 0) == 0)
        empty_rate = (empty_outputs / len(ok_rows)) if ok_rows else None

        def mean_num(key: str) -> Optional[float]:
            vals = [r[key] for r in ok_rows if isinstance(r.get(key), (int, float))]
            return statistics.mean(vals) if vals else None

        ws_sum.append([
            dep,
            len(dep_rows),
            len(ok_rows),
            len(dep_rows) - len(ok_rows),
            empty_outputs,
            empty_rate,
            statistics.mean(lat) if lat else None,
            statistics.median(lat) if lat else None,
            percentile(lat, 0.90) if lat else None,
            mean_num("prompt_tokens"),
            mean_num("completion_tokens"),
            mean_num("reasoning_tokens"),
            mean_num("total_tokens"),
        ])

    for col_idx, header in enumerate(sum_headers, start=1):
        ws_sum.column_dimensions[get_column_letter(col_idx)].width = max(18, len(header) + 2)

    wb.save(out_path)


def main() -> int:
    load_dotenv()

    endpoint = must_get("AZURE_OPENAI_ENDPOINT").rstrip("/")
    api_version = os.getenv("AZURE_OPENAI_API_VERSION", "2024-10-21")
    api_key = must_get("AZURE_OPENAI_API_KEY")

    deployments = parse_deployments()

    runs_per_deployment = int(os.getenv("RUNS_PER_DEPLOYMENT", "10"))
    temperature = float(os.getenv("TEMPERATURE", "0.2"))
    prompt = os.getenv(
        "TEST_PROMPT",
        "Tell me ONE short developer joke. Return ONLY the joke, one sentence, max 20 words. No extra text."
    )

    # Print answers? set PRINT_ANSWERS=1 in env to print full text
    print_answers = os.getenv("PRINT_ANSWERS", "0") == "1"

    print(f"Endpoint:    {endpoint}")
    print(f"API ver:     {api_version}")
    print(f"Deployments: {', '.join(deployments)}")
    print(f"Runs each:   {runs_per_deployment}")
    print(f"Temp:        {temperature}")
    print(f"Prompt:      {prompt}")

    client = AzureOpenAI(
        api_key=api_key,
        azure_endpoint=endpoint,
        api_version=api_version,
    )

    messages = [{"role": "user", "content": prompt}]
    raw_rows: List[Dict[str, Any]] = []

    for dep in deployments:
        print(f"\n=== {dep} ===")
        for i in range(1, runs_per_deployment + 1):
            t0 = time.perf_counter()
            ts = datetime.now(timezone.utc).isoformat()

            row: Dict[str, Any] = {
                "timestamp_utc": ts,
                "deployment": dep,
                "trial": i,
                "ok": False,
                "mode": None,
                "latency_ms": None,
                "prompt_tokens": None,
                "completion_tokens": None,
                "reasoning_tokens": None,
                "total_tokens": None,
                "finish_reason": None,
                "visible_chars": None,
                "error_http_status": None,
                "error_message": None,
                "output_preview": None,
            }

            try:
                resp, mode = call_chat_with_temp_fallback(
                    client=client,
                    deployment=dep,
                    messages=messages,
                    temperature=temperature,
                )

                dt_ms = (time.perf_counter() - t0) * 1000.0
                usage = safe_attr(resp, "usage")

                content = extract_text(resp) or ""
                preview = content.replace("\n", " ").strip()

                row.update({
                    "ok": True,
                    "mode": mode,
                    "latency_ms": round(dt_ms, 1),
                    "prompt_tokens": safe_attr(usage, "prompt_tokens"),
                    "completion_tokens": safe_attr(usage, "completion_tokens"),
                    "reasoning_tokens": extract_reasoning_tokens(resp),
                    "total_tokens": safe_attr(usage, "total_tokens"),
                    "finish_reason": safe_attr(resp.choices[0], "finish_reason"),
                    "visible_chars": len(content),
                    "output_preview": f"[{mode}] " + (preview[:200] + ("…" if len(preview) > 200 else "")),
                })

                # Metrics line
                print(
                    f"  run {i:02d}: {row['latency_ms']} ms | "
                    f"tok {row['prompt_tokens']}/{row['completion_tokens']} | "
                    f"reason={row['reasoning_tokens']} | {mode} | chars={row['visible_chars']}"
                )

                # ✅ Print answer (short preview always; full if PRINT_ANSWERS=1)
                if preview:
                    if print_answers:
                        print(f"       answer: {content.strip()}")
                    else:
                        print(f"       answer: {preview[:120]}{'…' if len(preview) > 120 else ''}")
                else:
                    print("       answer: <EMPTY>")

            except APIStatusError as e:
                dt_ms = (time.perf_counter() - t0) * 1000.0
                row["latency_ms"] = round(dt_ms, 1)
                row["error_http_status"] = e.status_code
                row["error_message"] = getattr(e.response, "text", "") or str(e)

                print(f"  run {i:02d}: FAIL HTTP {e.status_code} ({row['latency_ms']} ms)")
                print(f"           -> {row['error_message'][:400]}")

            except Exception as e:
                dt_ms = (time.perf_counter() - t0) * 1000.0
                row["latency_ms"] = round(dt_ms, 1)
                row["error_message"] = repr(e)

                print(f"  run {i:02d}: FAIL {repr(e)} ({row['latency_ms']} ms)")

            raw_rows.append(row)

    out_name = f"model_comparison_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%SZ')}.xlsx"
    out_path = os.path.abspath(out_name)
    write_excel(raw_rows, out_path)
    print(f"\n✅ Wrote results to: {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
